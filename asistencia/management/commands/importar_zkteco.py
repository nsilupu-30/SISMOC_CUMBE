"""Carga auditable y comprobación sin persistencia de las exportaciones originales."""
from collections import Counter
from contextlib import closing
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.core.exceptions import ValidationError
from django.db import transaction
import openpyxl

from asistencia.models import IdentidadBiometrica
from asistencia.services.zkteco_parser import ZKTecoExcelParser, normalizar


class Command(BaseCommand):
    help = "Importa marcaciones por sede y código. --simular revierte todos los cambios."

    def add_arguments(self, parser):
        parser.add_argument("archivo")
        parser.add_argument("--simular", action="store_true")
        parser.add_argument("--referencias", help="Libro revisado con hojas individuales para nombrar los códigos locales cortos.")

    def handle(self, *args, **options):
        archivo = Path(options["archivo"])
        if not archivo.is_file():
            raise CommandError("No se encontró el archivo de marcaciones.")
        try:
            with transaction.atomic():
                resultado = ZKTecoExcelParser(str(archivo), archivo.name).procesar()
                nombradas = 0
                if options["referencias"]:
                    referencias = Path(options["referencias"])
                    if not referencias.is_file():
                        raise CommandError("No se encontró el libro de referencia.")
                    sede = None
                    with closing(openpyxl.load_workbook(referencias, read_only=True, data_only=True)) as libro:
                        for hoja in libro:
                            if normalizar(hoja.cell(1, 4).value) == "fechahora":
                                sede = hoja.title.strip().upper()
                                continue
                            if sede is None:
                                continue
                            codigos = Counter()
                            for fila in hoja.iter_rows(max_col=4):
                                for idx in (2, 3):
                                    if isinstance(fila[idx].value, datetime):
                                        anterior = str(fila[idx - 1].value or "").strip()
                                        if anterior.isdigit() and len(anterior) < 8:
                                            codigos[anterior] += 1
                            if len(codigos) == 1:
                                codigo, cuenta = codigos.most_common(1)[0]
                                if cuenta >= 2:
                                    qs = IdentidadBiometrica.objects.filter(sede=sede, codigo=codigo, usuario__isnull=True)
                                    for identidad in qs:
                                        if not identidad.nombre or identidad.nombre == codigo:
                                            identidad.nombre = hoja.title.strip()
                                            identidad.save(update_fields=["nombre"])
                                            nombradas += 1
                self.stdout.write(f"Filas leídas: {resultado['filas_procesadas']}; nuevas: {resultado['marcaciones_nuevas']}; repetidas: {resultado['marcaciones_repetidas']}; códigos: {resultado['codigos_biometricos']}; pendientes de cuenta: {resultado['identidades_pendientes']}; nombres recuperados de referencia: {nombradas}.")
                if options["simular"]:
                    transaction.set_rollback(True)
                    self.stdout.write("Simulación finalizada: no se guardó ningún cambio.")
                else:
                    self.stdout.write(self.style.SUCCESS(f"Importación guardada. Lote #{resultado['carga_id']}. No se inventaron cuentas ni horas de trabajo."))
        except ValidationError as exc:
            raise CommandError(" ".join(exc.messages)) from exc
